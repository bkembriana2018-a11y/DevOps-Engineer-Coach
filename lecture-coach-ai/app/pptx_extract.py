from __future__ import annotations

from typing import Any

from pptx import Presentation


def extract_slides(file_bytes: bytes) -> list[dict[str, Any]]:
    """Pull title, body text, and speaker notes off every slide of a .pptx."""
    import io

    prs = Presentation(io.BytesIO(file_bytes))
    slides: list[dict[str, Any]] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        bullets: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            is_title = shape == slide.shapes.title
            if is_title and not title:
                title = text
                continue
            for line in text.splitlines():
                line = line.strip()
                if line:
                    bullets.append(line)
        notes = ""
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        slides.append(
            {
                "number": i,
                "title": title,
                "bullets": bullets,
                "notes": notes,
            }
        )
    return slides


def slides_to_text(slides: list[dict[str, Any]]) -> str:
    """Flatten extracted slides into plain text for LLM grounding or fallback rendering."""
    chunks = []
    for s in slides:
        heading = s["title"] or f"Slide {s['number']}"
        piece = f"\n\n--- Slide {s['number']}: {heading} ---\n"
        piece += "\n".join(f"- {b}" for b in s["bullets"])
        if s["notes"]:
            piece += f"\n(speaker notes: {s['notes']})"
        chunks.append(piece)
    return "".join(chunks).strip()
